import re
import zipfile
from io import BytesIO
from pathlib import Path, PurePosixPath

from docx import Document
from fastapi import UploadFile
from pypdf import PdfReader

MAX_UPLOAD_BYTES = 5 * 1024 * 1024
MAX_PDF_PAGES = 100
MAX_DOCX_MEMBERS = 512
MAX_DOCX_UNCOMPRESSED_BYTES = 20 * 1024 * 1024
MAX_DOCX_MEMBER_BYTES = 10 * 1024 * 1024
MAX_DOCX_COMPRESSION_RATIO = 100
MAX_EXTRACTED_TEXT_CHARS = 2 * 1024 * 1024
MIN_EXTRACTED_TEXT_CHARS = 100
NO_TEXT_ERROR = (
    "We could not extract text from this file. Please upload a text-based PDF/DOCX "
    "or paste your resume text."
)

#: The declared media type each supported extension must arrive with. The
#: extension is what selects the parser, so a part whose declared type disagrees
#: with its extension is a contradiction rather than a preference — ".pdf"
#: announced as "text/plain" would still be handed to the PDF parser. No browser
#: produces that pairing; only a hand-built request does.
CONTENT_TYPE_FOR_EXTENSION = {
    ".pdf": "application/pdf",
    ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    ".txt": "text/plain",
}
ALLOWED_CONTENT_TYPES = set(CONTENT_TYPE_FOR_EXTENSION.values())
ALLOWED_EXTENSIONS = set(CONTENT_TYPE_FOR_EXTENSION)


class DocumentParserError(ValueError):
    pass


def _resource_limit_error() -> DocumentParserError:
    return DocumentParserError(
        "This document is too complex to process safely. Please upload a smaller PDF/DOCX or paste your resume text."
    )


def validate_upload_metadata(filename: str | None, content_type: str | None) -> None:
    """Check the declared filename and media type before any bytes are read.

    Separated from the content checks so the route can refuse an unsupported
    upload without streaming it first: rejecting an unsupported extension costs
    nothing, and doing it early keeps unsupported bytes out of memory entirely.

    A part that declares NO content type is refused rather than accepted. It used
    to skip the media-type check altogether (``if content_type and ...``), which
    made the allow-list opt-out: omitting the header was enough to be judged on
    the filename alone. Browsers always send one — the spec substitutes
    ``application/octet-stream`` when the OS reports no type — so only a
    hand-built request reaches this branch.
    """
    extension = Path(filename or "").suffix.lower()
    declared = (content_type or "").split(";", 1)[0].strip().lower()

    if extension not in ALLOWED_EXTENSIONS:
        raise DocumentParserError("Unsupported file extension. Upload a PDF, DOCX, or TXT file.")
    if not declared:
        raise DocumentParserError("Uploaded file is missing its content type. Upload a PDF, DOCX, or TXT file.")
    if declared not in ALLOWED_CONTENT_TYPES:
        raise DocumentParserError("Unsupported file type. Upload a PDF, DOCX, or TXT file.")
    if declared != CONTENT_TYPE_FOR_EXTENSION[extension]:
        raise DocumentParserError(
            "The file's content type does not match its extension. Upload a PDF, DOCX, or TXT file."
        )


def validate_uploaded_file(file: UploadFile, content: bytes | None = None) -> None:
    """Full validation for an upload that has already been read into memory.

    The route validates metadata up front and bounds the size while streaming, so
    both checks are already satisfied by the time it calls in. They stay here as
    this helper's own contract, for any caller that hands over a complete
    ``bytes`` body instead of streaming it.
    """
    validate_upload_metadata(file.filename, file.content_type)
    if content is not None:
        if not content:
            raise DocumentParserError("Uploaded file is empty.")
        if len(content) > MAX_UPLOAD_BYTES:
            raise DocumentParserError("Uploaded file is too large. The limit is 5MB.")


def extract_text_from_pdf(content: bytes) -> str:
    try:
        if not content.startswith(b"%PDF-"):
            raise DocumentParserError(NO_TEXT_ERROR)
        reader = PdfReader(BytesIO(content))
        if reader.is_encrypted or len(reader.pages) > MAX_PDF_PAGES:
            raise _resource_limit_error()
        parts: list[str] = []
        extracted_chars = 0
        for page in reader.pages:
            page_text = page.extract_text() or ""
            extracted_chars += len(page_text)
            if extracted_chars > MAX_EXTRACTED_TEXT_CHARS:
                raise _resource_limit_error()
            parts.append(page_text)
        text = "\n".join(parts)
    except DocumentParserError:
        raise
    except Exception as exc:  # noqa: BLE001 - untrusted parser failures share one safe API error
        raise DocumentParserError(NO_TEXT_ERROR) from exc
    return normalize_extracted_text(text)


def _validate_docx_archive(content: bytes) -> None:
    """Validate the OPC container before python-docx expands or parses XML."""
    if not content.startswith(b"PK"):
        raise DocumentParserError(NO_TEXT_ERROR)
    try:
        with zipfile.ZipFile(BytesIO(content)) as archive:
            members = archive.infolist()
            if len(members) > MAX_DOCX_MEMBERS:
                raise _resource_limit_error()

            names: set[str] = set()
            total_uncompressed = 0
            for member in members:
                normalized_name = member.filename.replace("\\", "/")
                path = PurePosixPath(normalized_name)
                if not path.parts or path.is_absolute() or ".." in path.parts or ":" in path.parts[0]:
                    raise DocumentParserError(NO_TEXT_ERROR)
                if member.flag_bits & 0x1:
                    raise DocumentParserError(NO_TEXT_ERROR)
                if member.is_dir():
                    continue

                names.add(normalized_name)
                total_uncompressed += member.file_size
                if (
                    member.file_size > MAX_DOCX_MEMBER_BYTES
                    or total_uncompressed > MAX_DOCX_UNCOMPRESSED_BYTES
                    or (
                        member.file_size > 0
                        and member.file_size / max(member.compress_size, 1) > MAX_DOCX_COMPRESSION_RATIO
                    )
                ):
                    raise _resource_limit_error()

            if "[Content_Types].xml" not in names or "word/document.xml" not in names:
                raise DocumentParserError(NO_TEXT_ERROR)
    except DocumentParserError:
        raise
    except Exception as exc:  # noqa: BLE001 - malformed ZIP metadata must not escape as a 500
        raise DocumentParserError(NO_TEXT_ERROR) from exc


def extract_text_from_docx(content: bytes) -> str:
    _validate_docx_archive(content)
    try:
        document = Document(BytesIO(content))
        paragraphs = [paragraph.text for paragraph in document.paragraphs]
        table_cells = [
            cell.text
            for table in document.tables
            for row in table.rows
            for cell in row.cells
            if cell.text
        ]
    except Exception as exc:  # noqa: BLE001 - python-docx/lxml failures share one safe API error
        raise DocumentParserError(NO_TEXT_ERROR) from exc
    text = "\n".join([*paragraphs, *table_cells])
    if len(text) > MAX_EXTRACTED_TEXT_CHARS:
        raise _resource_limit_error()
    return normalize_extracted_text(text)


# Characters that PDF extraction commonly mangles, mapped to clean equivalents.
_CHAR_REPLACEMENTS = {
    "‣": "•",  # normalize assorted bullet glyphs to one form
    "▪": "•",
    "◦": "•",
    "⁃": "•",
    "": "•",  # Symbol-font bullet from Word/PDF
    "‘": "'",  # smart quotes
    "’": "'",
    "“": '"',
    "”": '"',
    "–": "-",  # en dash
    "—": "-",  # em dash
    "−": "-",  # minus sign
    " ": " ",  # non-breaking space
    "​": "",  # zero-width space
    "‌": "",  # zero-width non-joiner
    "‍": "",  # zero-width joiner
    "﻿": "",  # BOM
    "�": "",  # replacement character (hidden artifacts)
    "ﬁ": "fi",  # ligatures
    "ﬂ": "fl",
}

# Standalone page-number / artifact lines like "Page 1 of 2" or a bare "2".
_PAGE_ARTIFACT_RE = re.compile(r"^(page\s+\d+(\s+of\s+\d+)?|\d{1,3})$", re.IGNORECASE)


def normalize_extracted_text(text: str) -> str:
    for source, target in _CHAR_REPLACEMENTS.items():
        text = text.replace(source, target)
    # Re-join words broken across line boundaries by hyphenation ("infer-\nence").
    text = re.sub(r"([A-Za-z])-\n([a-z])", r"\1\2", text)

    normalized_lines: list[str] = []
    for line in text.splitlines():
        collapsed = re.sub(r"[ \t]+", " ", line).strip()
        if not collapsed:
            normalized_lines.append("")
            continue
        if _PAGE_ARTIFACT_RE.match(collapsed):
            continue  # drop page-number artifacts
        normalized_lines.append(collapsed)

    normalized = "\n".join(normalized_lines)
    return re.sub(r"\n{3,}", "\n\n", normalized).strip()


def detect_document_kind(text: str, filename: str | None = None) -> str:
    lower_text = text.lower()
    lower_name = (filename or "").lower()
    if "linkedin.com/in/" in lower_text or "contact info" in lower_text and "linkedin" in lower_text:
        return "linkedin_pdf"
    if "linkedin" in lower_name:
        return "linkedin_pdf"
    if any(marker in lower_text for marker in ["experience", "education", "skills"]):
        return "resume"
    return "unknown"


def extract_uploaded_file_text(file: UploadFile, content: bytes) -> str:
    validate_uploaded_file(file, content)
    extension = Path(file.filename or "").suffix.lower()
    if extension == ".pdf":
        text = extract_text_from_pdf(content)
    elif extension == ".docx":
        text = extract_text_from_docx(content)
    elif extension == ".txt":
        try:
            text = normalize_extracted_text(content.decode("utf-8"))
        except UnicodeDecodeError as exc:
            raise DocumentParserError("Text files must use UTF-8 encoding.") from exc
    else:
        raise DocumentParserError("Unsupported file extension. Upload a PDF, DOCX, or TXT file.")
    if len(text) < MIN_EXTRACTED_TEXT_CHARS:
        raise DocumentParserError(NO_TEXT_ERROR)
    return text
